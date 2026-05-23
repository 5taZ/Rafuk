"""Smoke + regression tests for the leads export endpoint.

Covers both formats:
  * format=csv (default) — same shape as before, with a new
    ``round`` step that lands cleanly on the spreadsheet side.
  * format=xlsx — openpyxl-rendered workbook with a styled header,
    number formatting on money columns, and a percent-suffixed ROI
    column. We don't validate the styling itself (openpyxl handles
    it deterministically), just that the bytes are a valid xlsx
    workbook with the expected sheet name + header row.
"""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.middleware.telegram_auth import TelegramInitData


def fake_telegram_user() -> TelegramInitData:
    return TelegramInitData(user_id=987_654, first_name="ExportUser", raw={})


def _seed_lead(client: TestClient) -> None:
    response = client.post(
        "/api/v1/leads",
        json={
            "query": "iPhone 15 128GB",
            "ad_id": 50101,
            "title": "iPhone 15 чёрный, как новый",
            "link": "https://www.kufar.by/item/50101",
            "price_byn": 1799,
            "target_resale_byn": 1899,
            "status": "new",
            "source": "manual",
        },
    )
    assert response.status_code == 201, response.text


def test_export_csv_round_trip() -> None:
    from api.dependencies import get_telegram_user
    from api.main import create_app

    app = create_app()
    app.dependency_overrides[get_telegram_user] = fake_telegram_user

    with TestClient(app) as client:
        _seed_lead(client)
        response = client.get("/api/v1/leads/export?format=csv")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        body = response.content.decode("utf-8")
        # Header + at least one data row.
        lines = [line for line in body.splitlines() if line.strip()]
        assert len(lines) >= 2
        assert lines[0].startswith("id,query,title,link,price_byn,")
        assert "iPhone 15 128GB" in body
        assert "50101" in body


def test_export_xlsx_returns_valid_workbook() -> None:
    """The new format=xlsx path renders a parseable .xlsx workbook
    with the same column schema as the CSV. We open the bytes back
    with openpyxl and assert the header row + the first data row."""
    from api.dependencies import get_telegram_user
    from api.main import create_app

    app = create_app()
    app.dependency_overrides[get_telegram_user] = fake_telegram_user

    with TestClient(app) as client:
        _seed_lead(client)
        response = client.get("/api/v1/leads/export?format=xlsx")
        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "leads_export.xlsx" in response.headers.get(
            "content-disposition", ""
        )

    workbook = load_workbook(io.BytesIO(response.content), read_only=False)
    assert workbook.sheetnames == ["Leads"]
    sheet = workbook["Leads"]
    header = [cell.value for cell in sheet[1]]
    assert header[:5] == ["id", "query", "title", "link", "price_byn"]

    first_row = [cell.value for cell in sheet[2]]
    assert first_row[1] == "iPhone 15 128GB"
    assert first_row[2] == "iPhone 15 чёрный, как новый"
    # price_byn column should be a number (Excel-formatted), not a
    # string — that's the whole reason xlsx is preferable to CSV.
    assert isinstance(first_row[4], (int, float))


def test_export_xlsx_neutralises_spreadsheet_formulas() -> None:
    """PR-04: a lead title or query starting with ``=``/``+``/``-``/``@``
    must be stored as text in both export shapes. The CSV path
    already prefixed a single quote; the XLSX path previously wrote
    the raw value and Excel would interpret it as a formula on open.
    """
    from api.dependencies import get_telegram_user
    from api.main import create_app

    app = create_app()
    app.dependency_overrides[get_telegram_user] = fake_telegram_user

    malicious_payload = {
        "query": "@SUM(1,2)",
        "ad_id": 50202,
        "title": "=cmd|'/c calc.exe'!A1",
        "link": "https://www.kufar.by/item/50202",
        "price_byn": 100.0,
        "status": "new",
        "source": "+manual",
    }

    with TestClient(app) as client:
        resp = client.post("/api/v1/leads", json=malicious_payload)
        assert resp.status_code == 201, resp.text

        # XLSX path
        xlsx = client.get("/api/v1/leads/export?format=xlsx")
        assert xlsx.status_code == 200
        wb = load_workbook(io.BytesIO(xlsx.content), read_only=False)
        sheet = wb["Leads"]
        first_row = [cell.value for cell in sheet[2]]
        # PR-04: every formula-triggering free-text cell starts with
        # a single quote — Excel renders this as literal text.
        assert first_row[1] == "'@SUM(1,2)"
        assert first_row[2] == "'=cmd|'/c calc.exe'!A1"
        assert first_row[7] == "'+manual"

        # CSV path keeps the same shape (already protected pre-wave135,
        # the assertion locks the contract in).
        csv_resp = client.get("/api/v1/leads/export?format=csv")
        assert csv_resp.status_code == 200
        body = csv_resp.content.decode("utf-8")
        assert "'@SUM(1,2)" in body
        assert "'=cmd|" in body
        assert "'+manual" in body


def test_export_rejects_unknown_format() -> None:
    """An invalid `format=` query value must be rejected by FastAPI's
    Literal validator with 422 — not silently picked.

    TEST-M2 / BE-M12: was previously ``status_code in (400, 422)``
    because the handler had a redundant ``if fmt not in (...): 400``
    check on top of the Literal validation. Wave 12 removed the
    handler-side branch (it was unreachable), so now the response is
    deterministically 422 from Pydantic / FastAPI.
    """
    from api.dependencies import get_telegram_user
    from api.main import create_app

    app = create_app()
    app.dependency_overrides[get_telegram_user] = fake_telegram_user

    with TestClient(app) as client:
        response = client.get("/api/v1/leads/export?format=pdf")
        assert response.status_code == 422


def test_export_csv_incomplete_cost_basis_shows_empty_profit() -> None:
    """E-FIND-09: a lead with sold_price but NULL buy_price must NOT
    show inflated profit — profit and ROI columns should be empty."""
    import csv as csv_mod

    from api.dependencies import get_telegram_user
    from api.main import create_app

    app = create_app()
    app.dependency_overrides[get_telegram_user] = fake_telegram_user

    with TestClient(app) as client:
        # Create a lead, then mark it sold without setting buy_price.
        resp = client.post(
            "/api/v1/leads",
            json={
                "query": "test incomplete",
                "ad_id": 60001,
                "title": "Incomplete cost basis item",
                "link": "https://www.kufar.by/item/60001",
                "price_byn": 900,
                "status": "new",
                "source": "manual",
            },
        )
        assert resp.status_code == 201, resp.text
        lead_id = resp.json()["id"]
        version = resp.json()["version"]
        # Mark sold without buy_price.
        patch_resp = client.patch(
            f"/api/v1/leads/{lead_id}",
            json={"status": "sold", "sold_price_byn": 1000.0, "version": version},
        )
        assert patch_resp.status_code == 200, patch_resp.text

        export_resp = client.get("/api/v1/leads/export?format=csv")
        assert export_resp.status_code == 200

    body = export_resp.content.decode("utf-8")
    reader = csv_mod.DictReader(io.StringIO(body))
    rows = [r for r in reader if r["id"] == str(lead_id)]
    assert len(rows) == 1
    row = rows[0]
    # Profit and ROI must be empty, not "1000.00".
    assert row["actual_profit"] == ""
    assert row["roi_percent"] == ""


def test_export_csv_sold_price_zero_reports_negative_profit() -> None:
    """LOGIC-NEW-1: sold_price_byn=0 is a legitimate giveaway — export must
    report negative profit (matching the API), not leave it blank."""
    import csv as csv_mod

    from api.dependencies import get_telegram_user
    from api.main import create_app

    app = create_app()
    app.dependency_overrides[get_telegram_user] = fake_telegram_user

    with TestClient(app) as client:
        resp = client.post(
            "/api/v1/leads",
            json={
                "query": "giveaway item",
                "ad_id": 70001,
                "title": "Free giveaway",
                "link": "https://www.kufar.by/item/70001",
                "price_byn": 100,
                "status": "new",
                "source": "manual",
            },
        )
        assert resp.status_code == 201
        lead_id = resp.json()["id"]
        version = resp.json()["version"]

        patch_resp = client.patch(
            f"/api/v1/leads/{lead_id}",
            json={
                "status": "sold",
                "buy_price_byn": 100,
                "sold_price_byn": 0,
                "version": version,
            },
        )
        assert patch_resp.status_code == 200

        export_resp = client.get("/api/v1/leads/export?format=csv")
        assert export_resp.status_code == 200

    body = export_resp.content.decode("utf-8")
    reader = csv_mod.DictReader(io.StringIO(body))
    rows = [r for r in reader if r["id"] == str(lead_id)]
    assert len(rows) == 1
    row = rows[0]
    # Profit must be -100.00 (0 - 100 - 0 expenses), not empty.
    assert row["actual_profit"] == "-100.00"


def test_export_csv_sold_price_none_reports_empty_profit() -> None:
    """LOGIC-NEW-1 regression: sold_price_byn=None must still produce
    empty profit — only the 0-case changed."""
    import csv as csv_mod

    from api.dependencies import get_telegram_user
    from api.main import create_app

    app = create_app()
    app.dependency_overrides[get_telegram_user] = fake_telegram_user

    with TestClient(app) as client:
        resp = client.post(
            "/api/v1/leads",
            json={
                "query": "unsold item",
                "ad_id": 70002,
                "title": "Still holding",
                "link": "https://www.kufar.by/item/70002",
                "price_byn": 200,
                "status": "bought",
                "source": "manual",
                "buy_price_byn": 200,
            },
        )
        assert resp.status_code == 201
        lead_id = resp.json()["id"]

        export_resp = client.get("/api/v1/leads/export?format=csv")
        assert export_resp.status_code == 200

    body = export_resp.content.decode("utf-8")
    reader = csv_mod.DictReader(io.StringIO(body))
    rows = [r for r in reader if r["id"] == str(lead_id)]
    assert len(rows) == 1
    row = rows[0]
    assert row["actual_profit"] == ""
    assert row["roi_percent"] == ""
