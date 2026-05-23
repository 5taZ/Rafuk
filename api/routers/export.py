from __future__ import annotations

import asyncio
import csv
import io
from collections import defaultdict
from typing import Any, Literal

from fastapi import APIRouter, Depends, Query, Request, Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from api.dependencies import get_session_factory_dependency, get_telegram_user
from api.limiter import limiter
from api.middleware.telegram_auth import TelegramInitData
from api.models import DealExpense, LeadItem
from api.services.workflow_store import resolve_user_id

router = APIRouter(tags=["export"])

CSV_HEADERS = [
    "id",
    "query",
    "title",
    "link",
    "price_byn",
    "target_resale_byn",
    "status",
    "source",
    "sold_price_byn",
    "sold_at",
    "total_expenses",
    "actual_profit",
    "roi_percent",
    "created_at",
    "updated_at",
]

# CSV injection protection: cells starting with these characters are
# interpreted as formulas by Excel/LibreOffice. Prefix with a single
# quote (the standard mitigation) so the cell is treated as text.
# https://owasp.org/www-community/attacks/CSV_Injection
_CSV_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def _csv_safe(value: object) -> object:
    """Sanitize a value for CSV output to prevent CSV injection."""
    if not isinstance(value, str):
        return value
    if value and value[0] in _CSV_FORMULA_TRIGGERS:
        return "'" + value
    return value


# Free-text column indices in CSV_HEADERS / _row_for_lead output that
# should be guarded against spreadsheet formula injection. Both the
# CSV writer (Excel/LibreOffice/Numbers) and the XLSX writer (Excel
# specifically) interpret a leading ``=``/``+``/``-``/``@`` as a
# formula even when the source column is plain text. We share the
# list so a single audit point covers both export paths (PR-04).
_SPREADSHEET_FREE_TEXT_COLUMNS = (1, 2, 3, 6, 7)


def _sanitize_row_for_spreadsheet(row: list[object]) -> list[object]:
    """Apply ``_csv_safe`` to every free-text column in ``row``.

    PR-04: previously only the CSV path called ``_csv_safe`` inline;
    the XLSX writer wrote raw cell values, so a Kufar title or
    user-typed note starting with ``=`` was interpreted as an Excel
    formula on first open. Wrap the row once, here, and reuse it
    from both writers so the two export shapes can't drift.
    """
    for idx in _SPREADSHEET_FREE_TEXT_COLUMNS:
        value = row[idx]
        if value is None:
            continue
        row[idx] = _csv_safe(value)
    return row


def _empty_csv_response() -> Response:
    output = io.StringIO()
    csv.writer(output).writerow(CSV_HEADERS)
    content = output.getvalue()
    output.close()
    return Response(
        content=content,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="leads_export.csv"'},
    )


def _row_for_lead(lead: LeadItem, lead_expenses: dict[int, float]) -> list[Any]:
    """Build the per-lead row payload shared by the CSV and XLSX writers."""
    # E-FIND-09 follow-up: when buy_price_byn is NULL but sold_price exists,
    # profit/ROI are indeterminate — output None instead of inflating.
    buy_price = float(lead.buy_price_byn) if lead.buy_price_byn is not None else None
    total_expenses = lead_expenses.get(lead.id, 0.0)

    # LOGIC-NEW-1: align with API — a recorded sold_price of 0 is a real giveaway, not "not sold".
    if lead.sold_price_byn is not None and buy_price is not None:
        sold_price = float(lead.sold_price_byn)
        total_cost = buy_price + total_expenses
        actual_profit = sold_price - total_cost
        roi_percent = (actual_profit / total_cost * 100) if total_cost > 0 else None
    else:
        actual_profit = None
        roi_percent = None

    return [
        lead.id,
        lead.query,
        lead.title,
        lead.link,
        float(lead.price_byn) if lead.price_byn is not None else None,
        float(lead.target_resale_byn) if lead.target_resale_byn is not None else None,
        lead.status,
        lead.source,
        float(lead.sold_price_byn) if lead.sold_price_byn is not None else None,
        lead.sold_at.isoformat() if lead.sold_at else None,
        round(total_expenses, 2),
        round(actual_profit, 2) if actual_profit is not None else None,
        round(roi_percent, 2) if roi_percent is not None else None,
        lead.created_at.isoformat(),
        lead.updated_at.isoformat(),
    ]


def _build_xlsx_workbook(leads: list[LeadItem], lead_expenses: dict[int, float]) -> bytes:
    """Render leads to a styled .xlsx workbook in memory.

    openpyxl is a runtime dep — added in pyproject.toml. Header is
    a bold accent row; profit / ROI columns get number formatting so
    the user can pivot in Excel without re-typing.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_align = Alignment(vertical="center")
    ws.append(CSV_HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 22

    # Numeric columns: price_byn, target_resale_byn, sold_price_byn,
    # total_expenses, actual_profit, roi_percent (1-based, post-header).
    numeric_cols = (5, 6, 9, 11, 12, 13)

    for lead in leads:
        # PR-04: sanitise free-text columns so a value starting with
        # ``=``/``+``/``-``/``@`` is treated as text, not as an Excel
        # formula. CSV path applies the same guard below.
        ws.append(_sanitize_row_for_spreadsheet(_row_for_lead(lead, lead_expenses)))

    # Format numbers — money to 2 decimals, ROI to 2 decimals + "%".
    for col in numeric_cols:
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col)
            if cell.value is None:
                continue
            cell.number_format = "0.00"
    # ROI specifically gets a percent suffix.
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=13)
        if cell.value is not None:
            cell.number_format = '0.00"%"'

    # Reasonable column widths so the export is usable on first open.
    widths = [6, 24, 36, 38, 12, 14, 14, 12, 14, 22, 14, 14, 10, 22, 22]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_FORMAT = Literal["csv", "xlsx"]


@router.get("/leads/export")
@limiter.limit("10/minute")
async def export_leads(
    request: Request,
    fmt: _FORMAT = Query("csv", alias="format"),
    telegram_user: TelegramInitData = Depends(get_telegram_user),
    session_factory: async_sessionmaker[AsyncSession] = Depends(get_session_factory_dependency),
) -> Response:
    """Export leads to a file. Supports CSV (default) and XLSX.

    M17 known limitation: this endpoint materializes the full result set
    (up to 5000 leads + expenses) in memory before serialising. For CSV
    output, a ``StreamingResponse`` with a generator that yields one row
    at a time would reduce peak memory, but the XLSX path requires the
    full workbook in memory (openpyxl limitation). Refactoring to stream
    CSV while keeping XLSX materialised is feasible but requires
    splitting the response path earlier; tracked for a future iteration.
    """
    # BE-M12: the explicit re-validation that lived here was unreachable —
    # ``fmt: _FORMAT`` (Literal["csv", "xlsx"]) is enforced by FastAPI's
    # query validation before the handler runs, so any non-allowed value
    # already returns 422 from the framework. Removing the dead branch.

    async with session_factory() as session:
        user_id = await resolve_user_id(session, telegram_user.user_id)
        if user_id is None:
            if fmt == "csv":
                return _empty_csv_response()
            # Empty xlsx fallback — header only, still a valid workbook.
            content = await asyncio.to_thread(_build_xlsx_workbook, [], {})
            return Response(
                content=content,
                media_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                headers={
                    "Content-Disposition": (
                        'attachment; filename="leads_export.xlsx"'
                    )
                },
            )

        result = await session.execute(
            select(LeadItem)
            .where(LeadItem.user_id == user_id)
            .order_by(LeadItem.created_at.desc())
            .limit(5000)
        )
        leads = list(result.scalars())

        lead_ids = [lead.id for lead in leads]
        lead_expenses: dict[int, float] = defaultdict(float)
        if lead_ids:
            expenses_result = await session.execute(
                select(DealExpense.lead_id, DealExpense.amount_byn).where(
                    DealExpense.lead_id.in_(lead_ids),
                    DealExpense.user_id == user_id,
                )
            )
            for lid, amount in expenses_result:
                lead_expenses[lid] += float(amount)

    if fmt == "csv":
        if not leads:
            return _empty_csv_response()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(CSV_HEADERS)
        for lead in leads:
            # PR-04: single sanitisation helper shared with the XLSX
            # writer above so the two export shapes can't drift.
            row = _sanitize_row_for_spreadsheet(_row_for_lead(lead, lead_expenses))
            # Per-cell None-to-empty + float formatting (CSV-only).
            row = [
                "" if value is None
                else (
                    f"{value:.2f}" if isinstance(value, float) else value
                )
                for value in row
            ]
            writer.writerow(row)
        content = output.getvalue()
        output.close()
        return Response(
            content=content,
            media_type="text/csv",
            headers={
                "Content-Disposition": 'attachment; filename="leads_export.csv"'
            },
        )

    # xlsx
    content_bytes = await asyncio.to_thread(_build_xlsx_workbook, leads, lead_expenses)
    return Response(
        content=content_bytes,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": 'attachment; filename="leads_export.xlsx"'
        },
    )
